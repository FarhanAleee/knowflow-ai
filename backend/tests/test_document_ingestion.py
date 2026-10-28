from pathlib import Path

import fitz
from docx import Document

from app.services.document_ingestion import extract_document, extract_pdf


def create_text_pdf(path: Path, page_texts: list[str]) -> Path:
    document = fitz.open()
    for text in page_texts:
        page = document.new_page()
        page.insert_text((72, 72), text)
    document.save(path)
    document.close()
    return path


def create_empty_pdf(path: Path) -> Path:
    pdf = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [] /Count 0 >>",
    ]
    for number, content in enumerate(objects, start=1):
        offsets.append(len(pdf))
        pdf.extend(f"{number} 0 obj\n".encode())
        pdf.extend(content)
        pdf.extend(b"\nendobj\n")

    xref_offset = len(pdf)
    pdf.extend(b"xref\n0 3\n0000000000 65535 f \n")
    for offset in offsets[1:]:
        pdf.extend(f"{offset:010} 00000 n \n".encode())
    pdf.extend(
        f"trailer\n<< /Size 3 /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF\n".encode()
    )
    path.write_bytes(pdf)
    return path


def test_extracts_text_and_page_metadata_from_generated_pdf(tmp_path):
    pdf_path = create_text_pdf(tmp_path / "team-notes.pdf", ["First page", "Second page"])

    result = extract_pdf(pdf_path)

    assert result.filename == "team-notes.pdf"
    assert result.total_pages == 2
    assert result.status == "success"
    assert [page.page_number for page in result.pages] == [1, 2]
    assert [page.text.strip() for page in result.pages] == ["First page", "Second page"]
    assert all(page.character_count == len(page.text) for page in result.pages)
    assert all(page.status == "extracted" for page in result.pages)


def test_reports_pages_without_extractable_text(tmp_path):
    pdf_path = tmp_path / "image-only.pdf"
    document = fitz.open()
    page = document.new_page()
    pixmap = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 10, 10), False)
    pixmap.clear_with(255)
    page.insert_image(page.rect, stream=pixmap.tobytes("png"))
    document.save(pdf_path)
    document.close()

    result = extract_pdf(pdf_path)

    assert result.status == "success"
    assert result.total_pages == 1
    assert result.pages[0].page_number == 1
    assert result.pages[0].text == ""
    assert result.pages[0].character_count == 0
    assert result.pages[0].status == "no_text"


def test_reports_empty_pdf(tmp_path):
    pdf_path = create_empty_pdf(tmp_path / "empty.pdf")

    result = extract_pdf(pdf_path)

    assert result.filename == "empty.pdf"
    assert result.total_pages == 0
    assert result.pages == []
    assert result.status == "empty"
    assert result.error == "PDF contains no pages."


def test_reports_invalid_pdf_without_raising(tmp_path):
    pdf_path = tmp_path / "invalid.pdf"
    pdf_path.write_bytes(b"This is not a PDF")

    result = extract_pdf(pdf_path)

    assert result.filename == "invalid.pdf"
    assert result.total_pages == 0
    assert result.pages == []
    assert result.status == "error"
    assert result.error == "Unable to open PDF file."


def test_unified_extraction_routes_pdf(tmp_path):
    pdf_path = create_text_pdf(tmp_path / "routed.pdf", ["PDF content"])

    result = extract_document(pdf_path)

    assert result.filename == "routed.pdf"
    assert result.content.strip() == "PDF content"
    assert result.total_pages == 1
    assert result.character_count == len(result.content)
    assert result.pages[0].page_number == 1


def test_extracts_docx_paragraphs_and_tables_in_document_order(tmp_path):
    docx_path = tmp_path / "team-notes.docx"
    document = Document()
    document.add_paragraph("Before table")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "First cell"
    table.cell(0, 1).text = "Second cell"
    document.add_paragraph("After table")
    document.save(docx_path)

    result = extract_document(docx_path)

    assert result.filename == "team-notes.docx"
    assert result.status == "success"
    assert result.total_pages is None
    assert result.pages == []
    assert [section.section_type for section in result.sections] == [
        "paragraph",
        "table",
        "paragraph",
    ]
    assert [section.section_number for section in result.sections] == [1, 2, 3]
    assert [section.page_number for section in result.sections] == [None, None, None]
    assert [section.text for section in result.sections] == [
        "Before table",
        "First cell | Second cell",
        "After table",
    ]
    assert result.content == "Before table\nFirst cell | Second cell\nAfter table"
    assert result.character_count == len(result.content)
    assert all(section.character_count == len(section.text) for section in result.sections)


def test_reports_corrupt_docx_without_raising(tmp_path):
    docx_path = tmp_path / "corrupt.docx"
    docx_path.write_bytes(b"not a DOCX package")

    result = extract_document(docx_path)

    assert result.filename == "corrupt.docx"
    assert result.total_pages is None
    assert result.sections == []
    assert result.status == "error"
    assert result.error == "Unable to read DOCX file."


def test_extracts_utf8_txt_without_physical_page_numbers(tmp_path):
    txt_path = tmp_path / "notes.txt"
    txt_path.write_text("Hello, knowledge!", encoding="utf-8")

    result = extract_document(txt_path)

    assert result.filename == "notes.txt"
    assert result.content == "Hello, knowledge!"
    assert result.character_count == len(result.content)
    assert result.total_pages is None
    assert result.pages == []
    assert len(result.sections) == 1
    assert result.sections[0].section_type == "text"
    assert result.sections[0].page_number is None


def test_extracts_utf8_bom_txt(tmp_path):
    txt_path = tmp_path / "bom.txt"
    txt_path.write_bytes(b"\xef\xbb\xbfHello")

    result = extract_document(txt_path)

    assert result.content == "Hello"
    assert result.character_count == 5
    assert result.status == "success"


def test_reports_invalid_utf8_txt_without_raising(tmp_path):
    txt_path = tmp_path / "invalid.txt"
    txt_path.write_bytes(b"\xff\xfe\xfa")

    result = extract_document(txt_path)

    assert result.content == ""
    assert result.character_count == 0
    assert result.status == "error"
    assert result.error == "Unable to read text file as UTF-8."


def test_reports_empty_txt(tmp_path):
    txt_path = tmp_path / "empty.txt"
    txt_path.write_text("", encoding="utf-8")

    result = extract_document(txt_path)

    assert result.content == ""
    assert result.character_count == 0
    assert result.sections == []
    assert result.total_pages is None
    assert result.status == "empty"


def test_reports_unsupported_extension(tmp_path):
    file_path = tmp_path / "notes.md"
    file_path.write_text("Not supported", encoding="utf-8")

    result = extract_document(file_path)

    assert result.filename == "notes.md"
    assert result.status == "error"
    assert result.error == "Unsupported file extension: .md."
